from __future__ import annotations

import argparse
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from statistics import mean
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT.parents[2] / "Grow_Naturally_Kitchen_Inventory_v10_Matching_and_Valuation.xlsx"
DEFAULT_ZIP = Path.home() / "Downloads" / "Brakes+1-6.zip"

INVOICE_LINE_HEADERS = [
    "Supplier Code",
    "Invoice Number",
    "Invoice Date",
    "Document Type",
    "Supplier Product Code",
    "Product Description",
    "Pack Size",
    "Quantity",
    "Unit Price (£)",
    "VAT Rate",
    "Line Value (£)",
    "Matched Internal Product ID",
    "Source File",
]
INVOICE_SUMMARY_HEADERS = [
    "Invoice Number",
    "Invoice Date",
    "Document Type",
    "Invoice Total (£)",
    "Lines Extracted",
    "Source File",
    "Extraction Status",
]
CATALOGUE_HEADERS = [
    "Supplier Code",
    "Supplier Product Code",
    "Latest Product Description",
    "Latest Pack Size",
    "Latest Unit Price (£)",
    "Average Unit Price (£)",
    "Latest Purchase Date",
    "Purchase Line Count",
    "VAT Rate",
]

LINE_PATTERN = re.compile(
    r"^(?P<code>\d{3,8})\s+(?P<body>.+?)\s+"
    r"(?P<qty>-?\d+(?:\.\d+)?)\s+"
    r"(?P<unit>-?\d+(?:\.\d+)?)\s+"
    r"(?P<vat>\d+(?:\.\d+)?)%\s+"
    r"(?P<value>-?[\d,]+(?:\.\d+)?)$"
)
ITEM_START_PATTERN = re.compile(r"^\d{3,8}\s+")
HEADER_BOUNDARY_PREFIXES = (
    "VAT Rate %",
    "Customer Signature:",
    "Payment terms",
    "Errors and Omissions",
)


@dataclass(frozen=True)
class BrakesLine:
    invoice_number: str
    invoice_date: datetime
    document_type: str
    supplier_product_code: str
    product_description: str
    pack_size: str
    quantity: float
    unit_price: float
    vat_rate: float
    line_value: float
    source_file: str


@dataclass(frozen=True)
class ImportResult:
    pdf_files: int
    unreadable_files: int
    raw_lines: int
    imported_lines: int
    unique_product_codes: int


def parse_brakes_zip(zip_path: Path) -> tuple[list[BrakesLine], int, int]:
    raw_lines: list[BrakesLine] = []
    unreadable_files = 0
    pdf_files = 0

    with ZipFile(zip_path) as archive:
        for name in archive.namelist():
            if not name.lower().endswith(".pdf"):
                continue
            pdf_files += 1
            data = archive.read(name)
            if not data:
                unreadable_files += 1
                continue
            try:
                raw_lines.extend(parse_brakes_pdf(name, data))
            except Exception:
                unreadable_files += 1

    return raw_lines, pdf_files, unreadable_files


def parse_brakes_pdf(source_file: str, pdf_data: bytes) -> list[BrakesLine]:
    reader = PdfReader(BytesIO(pdf_data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    text_lines = [line.strip() for line in text.splitlines() if line.strip()]

    invoice_number = next_value_after(text_lines, "Document Number")
    document_type = next_value_after(text_lines, "Document Type")
    date_text = next_value_after(text_lines, "Invoice Date") or next_value_after(text_lines, "Credit Date")
    invoice_date = parse_brakes_date(date_text)

    lines: list[BrakesLine] = []
    buffer = ""
    for text_line in text_lines:
        if any(text_line.startswith(prefix) for prefix in HEADER_BOUNDARY_PREFIXES):
            buffer = ""
            continue

        if ITEM_START_PATTERN.match(text_line):
            buffer = text_line
        elif buffer:
            buffer = f"{buffer} {text_line}"
        else:
            continue

        match = LINE_PATTERN.match(buffer)
        if not match:
            continue

        description, pack_size = split_product_body(match.group("body"))
        lines.append(
            BrakesLine(
                invoice_number=invoice_number,
                invoice_date=invoice_date,
                document_type=document_type,
                supplier_product_code=match.group("code"),
                product_description=description,
                pack_size=pack_size,
                quantity=float(match.group("qty")),
                unit_price=float(match.group("unit")),
                vat_rate=float(match.group("vat")) / 100,
                line_value=float(match.group("value").replace(",", "")),
                source_file=source_file,
            )
        )
        buffer = ""

    return lines


def next_value_after(lines: list[str], label: str) -> str:
    for index, line in enumerate(lines):
        if line == label and index + 1 < len(lines):
            return lines[index + 1].strip()
    return ""


def parse_brakes_date(value: str) -> datetime:
    return datetime.strptime(value, "%d.%m.%Y")


def split_product_body(body: str) -> tuple[str, str]:
    parts = re.split(r"\s{2,}", body.strip())
    if len(parts) >= 2:
        return " ".join(parts[:-1]).strip(), parts[-1].strip()

    fallback = re.match(
        r"^(?P<name>.+?)\s+(?P<pack>"
        r"(?:\d+\s*x\s*)?\d+(?:\.\d+)?(?:-\d+)?\s*(?:kg|g|ml|ltr|lt|oz|cm|cl)|"
        r"(?:\d+\s*x\s*)?\d+x\d+.*|"
        r"Each|\d+\s*Pack|WHTWINE|REDWINE|ROSEWINE|Ind"
        r")$",
        body.strip(),
        re.IGNORECASE,
    )
    if fallback:
        return fallback.group("name").strip(), fallback.group("pack").strip()
    return body.strip(), ""


def dedupe_lines(lines: list[BrakesLine]) -> list[BrakesLine]:
    deduped: dict[tuple[object, ...], BrakesLine] = {}
    for line in lines:
        key = (
            line.invoice_number,
            line.invoice_date,
            line.document_type,
            line.supplier_product_code,
            line.product_description,
            line.pack_size,
            line.quantity,
            line.unit_price,
            line.vat_rate,
            line.line_value,
        )
        deduped.setdefault(key, line)
    return list(deduped.values())


def write_brakes_workbook(workbook_path: Path, lines: list[BrakesLine]) -> None:
    workbook = load_workbook(workbook_path)
    write_invoice_lines(workbook["Invoice_Lines"], lines)
    write_invoice_summary(workbook["Invoice_Summary"], lines)
    write_catalogue(workbook["Brakes_Catalogue"], lines)
    workbook.save(workbook_path)


def write_invoice_lines(worksheet, lines: list[BrakesLine]) -> None:
    reset_sheet(worksheet, INVOICE_LINE_HEADERS)
    for line in sorted(lines, key=lambda item: (item.invoice_date, item.invoice_number, item.supplier_product_code)):
        worksheet.append(
            [
                "BRK",
                line.invoice_number,
                line.invoice_date,
                line.document_type,
                line.supplier_product_code,
                line.product_description,
                line.pack_size,
                line.quantity,
                line.unit_price,
                line.vat_rate,
                line.line_value,
                None,
                line.source_file,
            ]
        )


def write_invoice_summary(worksheet, lines: list[BrakesLine]) -> None:
    reset_sheet(worksheet, INVOICE_SUMMARY_HEADERS)
    grouped: dict[tuple[str, str], list[BrakesLine]] = defaultdict(list)
    for line in lines:
        grouped[(line.invoice_number, line.source_file)].append(line)

    for (_, source_file), invoice_lines in sorted(grouped.items(), key=lambda item: item[1][0].invoice_date):
        first = invoice_lines[0]
        invoice_total = round(sum(line.line_value * (1 + line.vat_rate) for line in invoice_lines), 2)
        worksheet.append(
            [
                first.invoice_number,
                first.invoice_date,
                first.document_type,
                invoice_total,
                len(invoice_lines),
                source_file,
                "Imported",
            ]
        )


def write_catalogue(worksheet, lines: list[BrakesLine]) -> None:
    reset_sheet(worksheet, CATALOGUE_HEADERS)
    by_code: dict[str, list[BrakesLine]] = defaultdict(list)
    for line in lines:
        by_code[line.supplier_product_code].append(line)

    for product_code in sorted(by_code, key=lambda code: int(code)):
        product_lines = by_code[product_code]
        latest = max(product_lines, key=lambda item: (item.invoice_date, item.invoice_number, item.source_file))
        worksheet.append(
            [
                "BRK",
                product_code,
                latest.product_description,
                latest.pack_size,
                latest.unit_price,
                round(mean(line.unit_price for line in product_lines), 4),
                latest.invoice_date,
                len(product_lines),
                latest.vat_rate,
            ]
        )


def reset_sheet(worksheet, headers: list[str]) -> None:
    worksheet.delete_rows(1, worksheet.max_row)
    worksheet.append(headers)


def import_brakes_invoices(zip_path: Path, workbook_path: Path, backup: bool = True) -> ImportResult:
    raw_lines, pdf_files, unreadable_files = parse_brakes_zip(zip_path)
    imported_lines = dedupe_lines(raw_lines)

    if backup:
        backup_path = workbook_path.with_suffix(f".before-brakes-import{workbook_path.suffix}")
        if not backup_path.exists():
            shutil.copy2(workbook_path, backup_path)

    write_brakes_workbook(workbook_path, imported_lines)
    return ImportResult(
        pdf_files=pdf_files,
        unreadable_files=unreadable_files,
        raw_lines=len(raw_lines),
        imported_lines=len(imported_lines),
        unique_product_codes=len({line.supplier_product_code for line in imported_lines}),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Brakes PDF invoices into the inventory workbook.")
    parser.add_argument("--zip", type=Path, default=DEFAULT_ZIP)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args()

    result = import_brakes_invoices(args.zip, args.workbook, backup=not args.no_backup)
    print(f"Read {result.pdf_files} Brakes PDF files")
    print(f"Skipped {result.unreadable_files} empty or unreadable PDF files")
    print(f"Parsed {result.raw_lines} raw invoice lines")
    print(f"Imported {result.imported_lines} de-duplicated invoice lines")
    print(f"Built {result.unique_product_codes} unique Brakes catalogue products")


if __name__ == "__main__":
    main()
