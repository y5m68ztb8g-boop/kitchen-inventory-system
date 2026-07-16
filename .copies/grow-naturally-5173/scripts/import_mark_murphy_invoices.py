from __future__ import annotations

import argparse
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from statistics import mean
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT.parents[2] / "Grow_Naturally_Kitchen_Inventory_v10_Matching_and_Valuation.xlsx"
DEFAULT_ZIP = Path.home() / "Downloads" / "Mark+Murphy+2-6.zip"

INVOICE_LINE_HEADERS = [
    "Supplier Code",
    "Invoice Number",
    "Invoice Date",
    "Document Type",
    "Supplier Product Code",
    "Product Description",
    "Pack Size",
    "Quantity",
    "Unit Price (£)",
    "VAT Rate",
    "Line Value (£)",
    "Matched Internal Product ID",
    "Source File",
]
INVOICE_SUMMARY_HEADERS = [
    "Invoice Number",
    "Invoice Date",
    "Document Type",
    "Invoice Total (£)",
    "Lines Extracted",
    "Source File",
    "Extraction Status",
]
CATALOGUE_HEADERS = [
    "Supplier Code",
    "Supplier Product Code",
    "Latest Product Description",
    "Latest Pack Size",
    "Latest Unit Price (£)",
    "Average Unit Price (£)",
    "Latest Purchase Date",
    "Purchase Line Count",
    "VAT Rate",
]

ITEM_START_PATTERN = re.compile(r"^\d{3,8}\s+")
OLD_LINE_PATTERN = re.compile(
    r"^(?P<code>\d{3,8})\s+(?P<body>.+?)\s+"
    r"(?P<qty>-?\d+(?:\.\d+)?)\s+"
    r"(?P<unit>-?\d+(?:\.\d+)?)\s+"
    r"(?P<vat>-?\d+(?:\.\d+)?)\s+"
    r"(?P<value>-?[\d,]+(?:\.\d+)?)$"
)
WEIGHT_LINE_PATTERN = re.compile(
    r"^(?P<code>\d{3,8})\s+(?P<body>.+?)\s+"
    r"(?P<weight>-?\d+(?:\.\d+)?)\s+"
    r"(?P<qty>-?\d+(?:\.\d+)?)\s+"
    r"(?P<unit>-?\d+(?:\.\d+)?)\s+"
    r"(?P<vat>-?\d+(?:\.\d+)?)\s+"
    r"(?P<value>-?[\d,]+(?:\.\d+)?)$"
)


@dataclass(frozen=True)
class MarkMurphyLine:
    invoice_number: str
    invoice_date: datetime
    document_type: str
    supplier_product_code: str
    product_description: str
    pack_size: str
    quantity: float
    unit_price: float
    vat_rate: float
    line_value: float
    source_file: str


@dataclass(frozen=True)
class MarkMurphySummary:
    invoice_number: str
    invoice_date: datetime | None
    document_type: str
    invoice_total: float | None
    lines_extracted: int
    source_file: str
    extraction_status: str


@dataclass(frozen=True)
class ImportResult:
    pdf_files: int
    imported_lines: int
    needs_review: int
    unique_product_codes: int


def parse_mark_murphy_zip(zip_path: Path) -> tuple[list[MarkMurphyLine], list[MarkMurphySummary], int]:
    lines: list[MarkMurphyLine] = []
    summaries: list[MarkMurphySummary] = []
    pdf_files = 0

    with ZipFile(zip_path) as archive:
        for name in archive.namelist():
            if not name.lower().endswith(".pdf"):
                continue
            pdf_files += 1
            data = archive.read(name)
            parsed_lines = parse_mark_murphy_pdf(name, data)
            lines.extend(parsed_lines)
            summaries.append(parse_mark_murphy_summary(name, data, parsed_lines))

    return dedupe_lines(lines), summaries, pdf_files


def parse_mark_murphy_pdf(source_file: str, pdf_data: bytes) -> list[MarkMurphyLine]:
    text_lines = extract_pdf_lines(pdf_data)
    document_type = "Statement" if "statement" in source_file.lower() else "Invoice"
    if document_type != "Invoice":
        return []

    invoice_number = find_invoice_number(text_lines, source_file)
    invoice_date = find_invoice_date(text_lines)
    uses_weight_column = any("CODE PRODUCT DESCRIPTION WGT QTY UNIT VAT GROSS" in line for line in text_lines)
    pattern = WEIGHT_LINE_PATTERN if uses_weight_column else OLD_LINE_PATTERN

    lines: list[MarkMurphyLine] = []
    in_items = False
    for text_line in text_lines:
        if text_line.startswith("CODE PRODUCT DESCRIPTION"):
            in_items = True
            continue
        if text_line.startswith("RATE VAT"):
            break
        if not in_items or not ITEM_START_PATTERN.match(text_line):
            continue

        match = pattern.match(text_line)
        if not match:
            continue

        lines.append(
            MarkMurphyLine(
                invoice_number=invoice_number,
                invoice_date=invoice_date,
                document_type=document_type,
                supplier_product_code=match.group("code"),
                product_description=clean_product_description(match.group("body")),
                pack_size="",
                quantity=float(match.group("qty")),
                unit_price=float(match.group("unit")),
                vat_rate=0,
                line_value=float(match.group("value").replace(",", "")),
                source_file=Path(source_file).name,
            )
        )

    return lines


def parse_mark_murphy_summary(
    source_file: str,
    pdf_data: bytes,
    parsed_lines: list[MarkMurphyLine],
) -> MarkMurphySummary:
    text_lines = extract_pdf_lines(pdf_data)
    document_type = "Statement" if "statement" in source_file.lower() else "Invoice"
    invoice_number = find_invoice_number(text_lines, source_file) if document_type == "Invoice" else ""
    invoice_date = find_invoice_date(text_lines) if document_type == "Invoice" else None
    invoice_total = find_invoice_total(text_lines)
    extraction_status = "Imported" if parsed_lines else "Needs review"

    return MarkMurphySummary(
        invoice_number=invoice_number,
        invoice_date=invoice_date,
        document_type=document_type,
        invoice_total=invoice_total,
        lines_extracted=len(parsed_lines),
        source_file=Path(source_file).name,
        extraction_status=extraction_status,
    )


def extract_pdf_lines(pdf_data: bytes) -> list[str]:
    reader = PdfReader(BytesIO(pdf_data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return [line.strip() for line in text.splitlines() if line.strip()]


def find_invoice_number(lines: list[str], source_file: str) -> str:
    filename_match = re.search(r"Invoice\s+(\d+)", source_file, re.IGNORECASE)
    if filename_match:
        return filename_match.group(1)

    for line in lines:
        match = re.search(r"\b(\d{7})\b", line)
        if match:
            return match.group(1)
    return ""


def find_invoice_date(lines: list[str]) -> datetime:
    for line in lines:
        match = re.match(r"^(\d{2}/\d{2}/\d{4})\s+00:\s+\d{7}$", line)
        if match:
            return datetime.strptime(match.group(1), "%d/%m/%Y")
    raise ValueError("Could not find Mark Murphy invoice date.")


def find_invoice_total(lines: list[str]) -> float | None:
    for line in lines:
        match = re.match(r"^Grand Total:\s+([\d,]+(?:\.\d+)?)$", line)
        if match:
            return float(match.group(1).replace(",", ""))
    return None


def clean_product_description(value: str) -> str:
    value = re.sub(r"\s+\.\s+\.\s+\.", " ", value)
    value = re.sub(r"\s+\.\s+", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip(" .")


def dedupe_lines(lines: list[MarkMurphyLine]) -> list[MarkMurphyLine]:
    deduped: dict[tuple[object, ...], MarkMurphyLine] = {}
    for line in lines:
        key = (
            line.invoice_number,
            line.invoice_date,
            line.supplier_product_code,
            line.product_description,
            line.quantity,
            line.unit_price,
            line.line_value,
        )
        deduped.setdefault(key, line)
    return list(deduped.values())


def write_mark_murphy_workbook(
    workbook_path: Path,
    lines: list[MarkMurphyLine],
    summaries: list[MarkMurphySummary],
) -> None:
    workbook = load_workbook(workbook_path)
    write_invoice_lines(workbook["MM_Invoice_Lines"], lines)
    write_invoice_summary(workbook["MM_Invoice_Summary"], summaries)
    write_catalogue(workbook["MM_Catalogue"], lines)
    update_import_progress(workbook, lines, summaries)
    workbook.save(workbook_path)


def write_invoice_lines(worksheet, lines: list[MarkMurphyLine]) -> None:
    reset_sheet(worksheet, INVOICE_LINE_HEADERS)
    for line in sorted(lines, key=lambda item: (item.invoice_date, item.invoice_number, item.supplier_product_code)):
        worksheet.append(
            [
                "MM",
                line.invoice_number,
                line.invoice_date,
                line.document_type,
                line.supplier_product_code,
                line.product_description,
                line.pack_size,
                line.quantity,
                line.unit_price,
                line.vat_rate,
                line.line_value,
                None,
                line.source_file,
            ]
        )


def write_invoice_summary(worksheet, summaries: list[MarkMurphySummary]) -> None:
    reset_sheet(worksheet, INVOICE_SUMMARY_HEADERS)
    for summary in sorted(summaries, key=lambda item: (item.invoice_date or datetime.min, item.source_file)):
        worksheet.append(
            [
                summary.invoice_number or None,
                summary.invoice_date,
                summary.document_type,
                summary.invoice_total,
                summary.lines_extracted,
                summary.source_file,
                summary.extraction_status,
            ]
        )


def write_catalogue(worksheet, lines: list[MarkMurphyLine]) -> None:
    reset_sheet(worksheet, CATALOGUE_HEADERS)
    grouped: dict[str, list[MarkMurphyLine]] = defaultdict(list)
    for line in lines:
        grouped[line.supplier_product_code].append(line)

    for product_code, product_lines in sorted(grouped.items()):
        latest = max(product_lines, key=lambda item: (item.invoice_date, item.invoice_number))
        prices = [line.unit_price for line in product_lines]
        worksheet.append(
            [
                "MM",
                product_code,
                latest.product_description,
                latest.pack_size,
                latest.unit_price,
                round(mean(prices), 4),
                latest.invoice_date,
                len(product_lines),
                latest.vat_rate,
            ]
        )


def update_import_progress(workbook, lines: list[MarkMurphyLine], summaries: list[MarkMurphySummary]) -> None:
    if "Import_Progress" not in workbook.sheetnames:
        return

    worksheet = workbook["Import_Progress"]
    for row in worksheet.iter_rows(min_row=2):
        if row[0].value == "Mark Murphy":
            row[1].value = len(summaries)
            row[2].value = len(lines)
            row[3].value = len({line.supplier_product_code for line in lines})
            row[4].value = sum(1 for summary in summaries if summary.extraction_status != "Imported")
            row[5].value = "Batch imported; statements need review"
            return


def reset_sheet(worksheet, headers: list[str]) -> None:
    worksheet.delete_rows(1, worksheet.max_row)
    worksheet.append(headers)


def import_mark_murphy_invoices(zip_path: Path, workbook_path: Path) -> ImportResult:
    lines, summaries, pdf_files = parse_mark_murphy_zip(zip_path)
    backup_path = workbook_path.with_suffix(".before-mm-import.xlsx")
    if not backup_path.exists():
        shutil.copy2(workbook_path, backup_path)
    write_mark_murphy_workbook(workbook_path, lines, summaries)
    return ImportResult(
        pdf_files=pdf_files,
        imported_lines=len(lines),
        needs_review=sum(1 for summary in summaries if summary.extraction_status != "Imported"),
        unique_product_codes=len({line.supplier_product_code for line in lines}),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Mark Murphy PDF invoices into the kitchen workbook.")
    parser.add_argument("--zip", type=Path, default=DEFAULT_ZIP)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()

    result = import_mark_murphy_invoices(args.zip, args.workbook)
    print(
        "Imported "
        f"{result.imported_lines} Mark Murphy lines from {result.pdf_files} PDFs; "
        f"{result.unique_product_codes} unique product codes; "
        f"{result.needs_review} documents need review."
    )


if __name__ == "__main__":
    main()
