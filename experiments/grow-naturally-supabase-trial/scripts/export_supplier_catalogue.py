from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT.parents[2] / "Grow_Naturally_Kitchen_Inventory_v10_Matching_and_Valuation.xlsx"
SUPPLIER_OUTPUT = ROOT / "src" / "generated" / "supplierCatalogue.ts"
FREEZER_OUTPUT = ROOT / "src" / "generated" / "freezerInventory.ts"

CATALOGUE_SHEETS = ["Brakes_Catalogue", "MM_Catalogue", "CMP_Catalogue"]
INVOICE_LINE_SHEETS = ["Invoice_Lines", "MM_Invoice_Lines", "CMP_Invoice_Lines"]
TRUSTED_MATCH_STATUSES = {"Exact supplier-code match", "High-confidence name match"}
MANUAL_SUPPLIER_PRODUCTS = [
    {
        "id": "BRK-135177",
        "supplierCode": "BRK",
        "supplierName": "Brakes / Sysco GB Ltd",
        "supplierProductCode": "135177",
        "productName": "Sysco Premium Coated Super Chunky Skin on Chips",
        "packSize": "4 x 2.5kg",
        "latestPrice": 20.74,
        "averagePrice": 20.74,
        "lowestPrice": 20.74,
        "highestPrice": 20.74,
        "latestPurchaseDate": "2026-07-08",
        "purchaseCount": 0,
        "vatRate": 0.2,
    }
]
MANUAL_FREEZER_MATCHES = {
    "BK004": ("BRK", "3625"),
    "BK006": ("BRK", "5011050"),
    "C004": ("BRK", "135096"),
    "C014": ("BRK", "146283"),
    "C016": ("BRK", "148602"),
    "D020": ("BRK", "136269"),
    "D021": ("BRK", "117350"),
    "D023": ("BRK", "123224"),
    "D028": ("BRK", "460806"),
    "FS005": ("CMP", "28HADFZIQF"),
    "FS008": ("CMP", "26HADSMO"),
    "PT002": ("BRK", "135177"),
}


def clean(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if value is None:
        return ""
    return value


def number(value):
    try:
        if value in ("", None):
            return None
        return round(float(value), 4)
    except (TypeError, ValueError):
        return None


def rows_by_header(workbook, sheet_name):
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    headers = [cell for cell in next(rows)]
    for row in rows:
        if any(cell not in ("", None) for cell in row):
            yield dict(zip(headers, row))


def build_freezer_inventory_entry(row):
    internal_id = str(row.get("Internal Product ID") or "").strip()
    product_name = str(row.get("Hotel Product Name") or "").strip()
    location_code = str(row.get("Location") or "").strip()
    quantity_text = str(row.get("Current Stock Text") or "").strip()
    if not internal_id or not product_name or not location_code:
        return None

    match_status = str(row.get("Match Status") or "").strip()
    is_trusted_match = match_status in TRUSTED_MATCH_STATUSES
    manual_match = MANUAL_FREEZER_MATCHES.get(internal_id)

    return {
        "id": internal_id,
        "productName": product_name,
        "locationCode": location_code,
        "quantityText": quantity_text,
        "recordedSupplierCode": str(row.get("Recorded Supplier Code") or "").strip(),
        "suggestedSupplierCode": manual_match[0]
        if manual_match
        else str(row.get("Matched Supplier") or "").strip()
        if is_trusted_match
        else "",
        "suggestedSupplierProductCode": manual_match[1]
        if manual_match
        else str(row.get("Matched Supplier Product Code") or "").strip()
        if is_trusted_match
        else "",
    }


def main():
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)

    suppliers = {}
    for row in rows_by_header(workbook, "Suppliers"):
        code = str(row.get("Supplier Code") or "").strip()
        if code:
            suppliers[code] = str(row.get("Supplier Name") or code).strip()

    invoice_stats = {}
    for sheet in INVOICE_LINE_SHEETS:
        for row in rows_by_header(workbook, sheet):
            supplier_code = str(row.get("Supplier Code") or "").strip()
            product_code = str(row.get("Supplier Product Code") or "").strip()
            price = number(row.get("Unit Price (£)"))
            if not supplier_code or not product_code or not price or price <= 0:
                continue
            key = (supplier_code, product_code)
            stats = invoice_stats.setdefault(key, {"prices": []})
            stats["prices"].append(price)

    entries = []
    for sheet in CATALOGUE_SHEETS:
        for row in rows_by_header(workbook, sheet):
            supplier_code = str(row.get("Supplier Code") or "").strip()
            product_code = str(row.get("Supplier Product Code") or "").strip()
            if not supplier_code or not product_code:
                continue

            prices = invoice_stats.get((supplier_code, product_code), {}).get("prices", [])
            latest_price = number(row.get("Latest Unit Price (£)"))
            average_price = number(row.get("Average Unit Price (£)"))

            entries.append(
                {
                    "id": f"{supplier_code}-{product_code}",
                    "supplierCode": supplier_code,
                    "supplierName": suppliers.get(supplier_code, supplier_code),
                    "supplierProductCode": product_code,
                    "productName": str(row.get("Latest Product Description") or "").strip(),
                    "packSize": str(row.get("Latest Pack Size") or "").strip(),
                    "latestPrice": latest_price or 0,
                    "averagePrice": average_price or (round(mean(prices), 4) if prices else 0),
                    "lowestPrice": min(prices) if prices else latest_price or 0,
                    "highestPrice": max(prices) if prices else latest_price or 0,
                    "latestPurchaseDate": clean(row.get("Latest Purchase Date")),
                    "purchaseCount": int(number(row.get("Purchase Line Count")) or len(prices) or 0),
                    "vatRate": number(row.get("VAT Rate")) or 0,
                }
            )

    existing_entry_ids = {entry["id"] for entry in entries}
    entries.extend(entry for entry in MANUAL_SUPPLIER_PRODUCTS if entry["id"] not in existing_entry_ids)
    entries.sort(key=lambda item: (item["supplierCode"], item["supplierProductCode"]))
    freezer_inventory = []
    for row in rows_by_header(workbook, "Product_Matching"):
        entry = build_freezer_inventory_entry(row)
        if entry:
            freezer_inventory.append(entry)

    SUPPLIER_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    SUPPLIER_OUTPUT.write_text(
        "import type { SupplierProduct } from \"../supplierProducts\";\n\n"
        "export const SUPPLIER_CATALOGUE: SupplierProduct[] = "
        + json.dumps(entries, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    FREEZER_OUTPUT.write_text(
        "export type FreezerInventoryItem = {\n"
        "  id: string;\n"
        "  productName: string;\n"
        "  locationCode: string;\n"
        "  quantityText: string;\n"
        "  recordedSupplierCode: string;\n"
        "  suggestedSupplierCode: string;\n"
        "  suggestedSupplierProductCode: string;\n"
        "};\n\n"
        "export const FREEZER_INVENTORY: FreezerInventoryItem[] = "
        + json.dumps(freezer_inventory, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    print(f"Exported {len(entries)} supplier products to {SUPPLIER_OUTPUT}")
    print(f"Exported {len(freezer_inventory)} freezer inventory items to {FREEZER_OUTPUT}")


if __name__ == "__main__":
    main()
